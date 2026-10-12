"""Import accounts from CSV or Excel with column mapping."""

from __future__ import annotations

import csv
from pathlib import Path

from PySide6.QtWidgets import (
    QComboBox,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
    QWizard,
    QWizardPage,
)

from flowpilot.storage.models import Account
from flowpilot.ui.app_context import AppContext
from flowpilot.ui.widgets.clay import ClayButton, muted_label

TARGETS = [("Ignore", "ignore"), ("Name", "name"), ("Profile ID", "remote_profile_id"), ("Group", "group"),
           ("Tags", "tags"), ("Notes", "notes"), ("Variable", "variable")]
GUESS = {"name": "name", "profile": "remote_profile_id", "profile_id": "remote_profile_id", "id": "remote_profile_id",
         "user_id": "remote_profile_id", "uuid": "remote_profile_id", "group": "group", "tags": "tags", "tag": "tags",
         "notes": "notes", "note": "notes", "remark": "notes"}


def read_table(path: Path) -> tuple[list[str], list[list[str]]]:
    """Read a CSV/XLSX file into (headers, rows)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raw = path.read_text("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(raw.splitlines(), dialect)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    body = [r + [""] * (len(headers) - len(r)) for r in rows[1:]]
    return headers, body


def guess_target(header: str) -> str:
    key = header.strip().lower().replace(" ", "_")
    return GUESS.get(key, "variable")


class _FilePage(QWizardPage):
    def __init__(self, wizard: ImportWizard) -> None:
        super().__init__()
        self.wiz = wizard
        self.setTitle("Choose a file")
        self.setSubTitle("A CSV or Excel file with one account per row. The first row must contain column names.")
        lay = QVBoxLayout(self)
        row = QHBoxLayout()
        self.path = QLineEdit()
        self.path.setReadOnly(True)
        browse = ClayButton("Browse…", "folder")
        browse.clicked.connect(self._browse)
        row.addWidget(self.path, 1)
        row.addWidget(browse)
        lay.addLayout(row)
        self.info = muted_label("")
        lay.addWidget(self.info)
        lay.addStretch(1)

    def _browse(self) -> None:
        file, _ = QFileDialog.getOpenFileName(self, "Choose file", "", "Tables (*.csv *.txt *.xlsx *.xlsm)")
        if not file:
            return
        try:
            self.wiz.headers, self.wiz.rows = read_table(Path(file))
        except Exception as exc:  # noqa: BLE001
            self.info.setText(f"Could not read the file: {exc}")
            return
        self.path.setText(file)
        self.info.setText(f"Found {len(self.wiz.rows)} row(s) and {len(self.wiz.headers)} column(s).")
        self.completeChanged.emit()

    def isComplete(self) -> bool:  # noqa: N802
        return bool(self.wiz.rows)


class _MapPage(QWizardPage):
    def __init__(self, wizard: ImportWizard) -> None:
        super().__init__()
        self.wiz = wizard
        self.setTitle("Match the columns")
        self.setSubTitle("Tell FlowPilot what each column means. Other columns can become variables "
                         "(e.g. a column 'email' becomes {{account.email}}).")
        lay = QVBoxLayout(self)
        conn_row = QHBoxLayout()
        conn_row.addWidget(QLabel("Connection for these accounts"))
        self.conn = QComboBox()
        conn_row.addWidget(self.conn, 1)
        lay.addLayout(conn_row)
        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["Column", "Means", "Variable name"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.verticalHeader().setVisible(False)
        lay.addWidget(self.table)

    def initializePage(self) -> None:  # noqa: N802
        self.conn.clear()
        for c in self.wiz.app.repos.connections.list():
            self.conn.addItem(c.name, c.id)
        self.table.setRowCount(0)
        for h in self.wiz.headers:
            r = self.table.rowCount()
            self.table.insertRow(r)
            item = QTableWidgetItem(h)
            self.table.setItem(r, 0, item)
            box = QComboBox()
            for text, value in TARGETS:
                box.addItem(text, value)
            box.setCurrentIndex(box.findData(guess_target(h)))
            self.table.setCellWidget(r, 1, box)
            self.table.setItem(r, 2, QTableWidgetItem(h.strip().lower().replace(" ", "_")))

    def mapping(self) -> list[tuple[str, str]]:
        out = []
        for r in range(self.table.rowCount()):
            box = self.table.cellWidget(r, 1)
            out.append((box.currentData(), self.table.item(r, 2).text().strip()))  # type: ignore[union-attr]
        return out

    def validatePage(self) -> bool:  # noqa: N802
        targets = [t for t, _ in self.mapping()]
        if "remote_profile_id" not in targets:
            self.setSubTitle("⚠ Choose which column contains the Profile ID.")
            return False
        return True


class _DonePage(QWizardPage):
    def __init__(self, wizard: ImportWizard) -> None:
        super().__init__()
        self.wiz = wizard
        self.setTitle("Import")
        lay = QVBoxLayout(self)
        self.label = muted_label("")
        lay.addWidget(self.label)

    def initializePage(self) -> None:  # noqa: N802
        created, updated = self.wiz.do_import()
        self.label.setText(f"Done: {created} new account(s), {updated} updated.")


class ImportWizard(QWizard):
    def __init__(self, app: AppContext, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.app = app
        self.setWindowTitle("Import accounts")
        self.resize(720, 560)
        self.headers: list[str] = []
        self.rows: list[list[str]] = []
        self.file_page = _FilePage(self)
        self.map_page = _MapPage(self)
        self.addPage(self.file_page)
        self.addPage(self.map_page)
        self.addPage(_DonePage(self))
        self.setOption(QWizard.WizardOption.NoBackButtonOnLastPage, True)

    def do_import(self) -> tuple[int, int]:
        repos = self.app.repos
        conn_id = self.map_page.conn.currentData()
        mapping = self.map_page.mapping()
        groups = {g.name: g for g in repos.groups.list()}
        created = updated = 0
        for row in self.rows:
            values: dict[str, str] = {}
            variables: dict[str, str] = {}
            for (target, var_name), cell in zip(mapping, row):
                cell = cell.strip()
                if target == "variable" and var_name and cell:
                    variables[var_name.removeprefix("account.")] = cell
                elif target not in ("ignore", "variable"):
                    values[target] = cell
            remote = values.get("remote_profile_id", "")
            if not remote:
                continue
            acc = repos.accounts.find_by_remote(conn_id, remote) if conn_id else None
            is_new = acc is None
            acc = acc or Account(connection_id=conn_id, remote_profile_id=remote)
            acc.name = values.get("name") or acc.name or remote
            if values.get("group"):
                g = groups.get(values["group"]) or repos.groups.create(values["group"])
                groups[g.name] = g
                acc.group_id = g.id
            if "tags" in values:
                acc.tags = values["tags"]
            if "notes" in values:
                acc.notes = values["notes"]
            merged = ({} if is_new else repos.accounts.variables(acc)) | variables
            repos.accounts.save(acc, variables=merged)
            created += is_new
            updated += not is_new
        self.app.bridge.data_changed.emit("accounts")
        self.app.bridge.data_changed.emit("groups")
        return created, updated
